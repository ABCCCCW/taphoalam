from __future__ import annotations

from datetime import datetime, timedelta

from fastapi import APIRouter, Depends, HTTPException, WebSocket, WebSocketDisconnect
from pydantic import BaseModel
from sqlalchemy.orm import Session, joinedload

from app.config import settings
from app.core.paging import clamp_page, page_meta
from app.core.permissions import require
from app.core.utils import money, next_code, parse_iso_date, utcnow
from app.services.promo_service import day_end_utc, day_start_utc, live_order_promos, serialize_promo
from app.services.batch_service import expiry_summary
from app.database import SessionLocal, get_db
from app.deps import get_current_staff
from app.models import AuditLog, Customer, Order, Promotion, Setting, Shift, User
from app.realtime.hub import hub

router = APIRouter(prefix="/api/v1", tags=["ops"])


@router.get("/customers")
def list_customers(
    phone: str | None = None,
    q: str | None = None,
    page: int = 1,
    size: int = 10,
    db: Session = Depends(get_db),
    _: User = Depends(require("customer.*")),
):
    page, size = clamp_page(page, size)
    query = db.query(Customer)
    if phone:
        query = query.filter(Customer.phone == phone)
    if q:
        query = query.filter((Customer.name.ilike(f"%{q}%")) | (Customer.phone.ilike(f"%{q}%")))
    total = query.count()
    rows = query.order_by(Customer.id.desc()).offset((page - 1) * size).limit(size).all()
    items = [
        {
            "id": c.id,
            "code": c.code,
            "name": c.name,
            "phone": c.phone,
            "email": c.email,
            "tier": c.tier,
            "loyalty_points": c.loyalty_points,
            "total_spent": float(c.total_spent),
            "address": c.address,
            "source": c.source,
        }
        for c in rows
    ]
    return {"items": items, **page_meta(page, size, total)}


class CustomerIn(BaseModel):
    name: str
    phone: str | None = None
    email: str | None = None
    address: str | None = None


@router.post("/customers")
def create_customer(body: CustomerIn, db: Session = Depends(get_db), _: User = Depends(require("customer.*"))):
    if body.phone and db.query(Customer).filter(Customer.phone == body.phone).first():
        raise HTTPException(409, "Số điện thoại đã tồn tại")
    seq = db.query(Customer).count() + 1
    c = Customer(code=f"KH{seq:06d}", name=body.name, phone=body.phone, email=body.email, address=body.address, source="POS")
    db.add(c)
    db.flush()
    return {"id": c.id, "code": c.code, "name": c.name, "phone": c.phone, "loyalty_points": 0, "tier": c.tier}


class ShiftOpenIn(BaseModel):
    opening_cash: float = 0


@router.post("/shifts/open")
def open_shift(body: ShiftOpenIn, db: Session = Depends(get_db), user: User = Depends(require("shift.own"))):
    existing = db.query(Shift).filter(Shift.user_id == user.id, Shift.status == "OPEN").first()
    if existing:
        return _shift(existing)
    day = utcnow().strftime("%Y%m%d")
    seq = db.query(Shift).filter(Shift.code.like(f"CA{day}-%")).count() + 1
    shift = Shift(
        code=next_code("CA", seq),
        user_id=user.id,
        warehouse_id=settings.default_warehouse_id,
        opening_cash=body.opening_cash,
        expected_cash=body.opening_cash,
    )
    db.add(shift)
    db.flush()
    return _shift(shift)


@router.get("/shifts/current")
def current_shift(db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    shift = db.query(Shift).filter(Shift.user_id == user.id, Shift.status == "OPEN").first()
    return _shift(shift) if shift else None


class ShiftCloseIn(BaseModel):
    closing_cash: float


@router.post("/shifts/{shift_id}/close")
def close_shift(shift_id: int, body: ShiftCloseIn, db: Session = Depends(get_db), user: User = Depends(require("shift.own"))):
    shift = db.get(Shift, shift_id)
    if not shift or shift.status != "OPEN":
        raise HTTPException(400, "Ca không hợp lệ")
    if user.role != "ADMIN" and shift.user_id != user.id:
        raise HTTPException(403, "Không đóng ca của người khác")
    shift.closing_cash = body.closing_cash
    shift.expected_cash = money(float(shift.opening_cash) + float(shift.cash_sales))
    shift.cash_diff = money(body.closing_cash - float(shift.expected_cash))
    shift.status = "CLOSED"
    shift.closed_at = utcnow()
    return _shift(shift)


def _shift(s: Shift):
    return {
        "id": s.id,
        "code": s.code,
        "status": s.status,
        "opening_cash": float(s.opening_cash),
        "cash_sales": float(s.cash_sales),
        "qr_sales": float(s.qr_sales),
        "expected_cash": float(s.expected_cash),
        "closing_cash": float(s.closing_cash) if s.closing_cash is not None else None,
        "cash_diff": float(s.cash_diff) if s.cash_diff is not None else None,
        "total_orders": s.total_orders,
        "opened_at": s.opened_at.isoformat() if s.opened_at else None,
        "cashier": s.user.full_name if s.user else None,
    }


@router.get("/reports/dashboard")
def dashboard(db: Session = Depends(get_db), user: User = Depends(require("product.read"))):
    if user.role == "CASHIER":
        raise HTTPException(403, "Thu ngân không xem được báo cáo")
    today = utcnow().date()
    orders = db.query(Order).filter(Order.status == "COMPLETED").all()
    today_orders = [o for o in orders if o.completed_at and o.completed_at.date() == today]
    yesterday_orders = [o for o in orders if o.completed_at and o.completed_at.date() == (today.fromordinal(today.toordinal() - 1))]
    revenue_today = sum(float(o.total_amount) for o in today_orders)
    revenue_yesterday = sum(float(o.total_amount) for o in yesterday_orders)
    profit_today = 0.0
    for o in today_orders:
        for i in o.items:
            profit_today += float(i.line_total) - float(i.cost_price) * float(i.quantity)
    from app.models import Inventory

    # Danh sách món sắp hết để Tổng quan hiện ngay, món cạn nhất (so với mức tối thiểu) lên đầu.
    low_items = []
    for inv in db.query(Inventory).options(joinedload(Inventory.product)).all():
        p = inv.product
        if p and p.is_active and float(inv.quantity) <= p.min_stock:
            low_items.append({
                "id": p.id,
                "name": p.name,
                "image_url": p.image_url,
                "emoji": p.emoji,
                "quantity": float(inv.quantity),
                "min_stock": p.min_stock,
            })
    low_items.sort(key=lambda x: (x["quantity"] / x["min_stock"]) if x["min_stock"] else x["quantity"])
    low = len(low_items)
    cash_amt = 0.0
    qr_amt = 0.0
    for o in today_orders:
        for p in o.payments:
            if p.status != "SUCCESS":
                continue
            if p.method == "CASH":
                cash_amt += float(p.amount)
            else:
                qr_amt += float(p.amount)
    online_rows = (
        db.query(Order)
        .filter(Order.channel == "ONLINE", Order.status.in_(["PENDING_PAYMENT", "PENDING_CONFIRM", "CONFIRMED", "PACKING"]))
        .order_by(Order.created_at)
        .all()
    )
    online_pending = len(online_rows)
    from app.models import Product

    cost_pending = db.query(Product).filter(Product.cost_confirmed.is_(False), Product.is_active.is_(True)).count()
    can_profit = user.role == "ADMIN"
    lots = expiry_summary(db)
    return {
        "revenue_today": revenue_today,
        "revenue_yesterday": revenue_yesterday,
        "orders_today": len(today_orders),
        "aov": money(revenue_today / len(today_orders)) if today_orders else 0,
        "profit_today": profit_today if can_profit else None,
        "low_stock": low,
        "low_stock_items": low_items,
        "online_items": [
            {
                "id": o.id,
                "code": o.code,
                "status": o.status,
                "customer_name": o.customer.name if o.customer else None,
                "total": float(o.total_amount),
                "created_at": o.created_at.isoformat() if o.created_at else None,
            }
            for o in online_rows
        ],
        "cash_amount": cash_amt,
        "qr_amount": qr_amt,
        "online_pending": online_pending,
        "cost_pending": cost_pending,
        "expired_count": lots["expired_count"],
        "expiring_count": lots["expiring_count"],
        "expired": lots["expired"],
        "expiring": lots["expiring"],
        "warn_days": lots["warn_days"],
    }


@router.get("/reports/revenue")
def revenue(group_by: str = "day", channel: str | None = None, db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị xem báo cáo doanh thu")
    query = db.query(Order).filter(Order.status == "COMPLETED")
    if channel:
        query = query.filter(Order.channel == channel)
    rows = query.all()
    from collections import defaultdict

    buckets = defaultdict(lambda: {"revenue": 0, "orders": 0, "pos": 0, "online": 0})
    for o in rows:
        if not o.completed_at:
            continue
        key = o.completed_at.date().isoformat()
        buckets[key]["revenue"] += float(o.total_amount)
        buckets[key]["orders"] += 1
        buckets[key][o.channel.lower()] += float(o.total_amount)
    out = [{"date": k, **v} for k, v in sorted(buckets.items())]
    return out


VN_OFFSET = timedelta(hours=7)


def _local(dt: datetime) -> datetime:
    """Giờ Hà Nội — đơn chốt lúc 1h sáng phải tính cho hôm nay, không phải hôm qua (UTC)."""
    return dt + VN_OFFSET


@router.get("/reports/summary")
def report_summary(days: int = 30, db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    """Mọi số liệu cho Báo cáo / Tổng quan trong một khoảng ngày, kèm kỳ liền trước để so."""
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị xem báo cáo")
    return build_summary(db, days)


@router.get("/reports/export")
def report_export(days: int = 30, db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    """Tải báo cáo doanh thu dạng Excel (.xlsx): tóm tắt, theo ngày, theo mặt hàng."""
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị xuất báo cáo")
    import io

    from fastapi.responses import StreamingResponse
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    r = build_summary(db, days)
    t, p = r["totals"], r["previous"]
    money_fmt = '#,##0" ₫"'
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="18181B")

    def table(ws, header, rows, money_cols=(), widths=None, start_row=1):
        for c, h in enumerate(header, 1):
            cell = ws.cell(row=start_row, column=c, value=h)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center")
        for i, row in enumerate(rows, start_row + 1):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=i, column=c, value=v)
                if c in money_cols:
                    cell.number_format = money_fmt
        for c, w in enumerate(widths or [], 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)

    def dmy(iso):
        y, m, d = iso.split("-")
        return f"{d}/{m}/{y}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Tóm tắt"
    ws["A1"] = "BÁO CÁO DOANH THU — LÂM LY MART"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Từ {dmy(r['from'])} đến {dmy(r['to'])} ({r['days']} ngày)"
    table(
        ws,
        ["Chỉ số", "Kỳ này", "Kỳ trước", "Thay đổi"],
        [
            ["Doanh thu", t["revenue"], p["revenue"], (t["revenue"] - p["revenue"]) / p["revenue"] if p["revenue"] else None],
            ["Lãi gộp", t["profit"], p["profit"], (t["profit"] - p["profit"]) / p["profit"] if p["profit"] else None],
            ["Số đơn", t["orders"], p["orders"], (t["orders"] - p["orders"]) / p["orders"] if p["orders"] else None],
            ["Trung bình mỗi đơn", t["aov"], None, None],
            ["Tỷ suất lãi", t["margin"] / 100, None, None],
            ["Tiền mặt", r["payments"]["CASH"], None, None],
            ["Chuyển khoản", r["payments"]["QR"] + r["payments"]["OTHER"], None, None],
        ],
        widths=[24, 18, 18, 12],
        start_row=4,
    )
    for row in range(5, 12):
        for col in (2, 3):
            if row not in (7, 9):
                ws.cell(row=row, column=col).number_format = money_fmt
        ws.cell(row=row, column=4).number_format = "0%"
    ws.cell(row=9, column=2).number_format = "0.0%"
    ws.freeze_panes = None

    ws2 = wb.create_sheet("Theo ngày")
    table(
        ws2,
        ["Ngày", "Doanh thu", "Quầy", "Website", "Lãi gộp", "Số đơn"],
        [[dmy(x["date"]), x["revenue"], x["pos"], x["online"], x["profit"], x["orders"]] for x in r["daily"]]
        + [["Tổng", t["revenue"], sum(x["pos"] for x in r["daily"]), sum(x["online"] for x in r["daily"]), t["profit"], t["orders"]]],
        money_cols=(2, 3, 4, 5),
        widths=[14, 16, 16, 16, 16, 10],
    )
    ws2.cell(row=len(r["daily"]) + 2, column=1).font = Font(bold=True)

    ws3 = wb.create_sheet("Mặt hàng")
    table(
        ws3,
        ["#", "Mặt hàng", "Đã bán", "Doanh thu", "Lãi gộp", "Tỷ suất"],
        [
            [i, x["name"], x["qty"], x["revenue"], x["profit"], (x["profit"] / x["revenue"]) if x["revenue"] else 0]
            for i, x in enumerate(r["products"], 1)
        ],
        money_cols=(4, 5),
        widths=[6, 40, 10, 16, 16, 10],
    )
    for row in range(2, len(r["products"]) + 2):
        ws3.cell(row=row, column=6).number_format = "0%"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"bao-cao-doanh-thu-{r['from']}-{r['to']}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


def build_summary(db: Session, days: int) -> dict:
    days = max(1, min(int(days), 366))
    from collections import defaultdict

    from app.core.utils import shop_today
    from app.models import Category, Product

    today = shop_today()
    start = today - timedelta(days=days - 1)
    prev_start = start - timedelta(days=days)
    since_utc = datetime.combine(prev_start, datetime.min.time()) - VN_OFFSET

    orders = (
        db.query(Order)
        .filter(Order.status == "COMPLETED", Order.completed_at >= since_utc)
        .options(joinedload(Order.items), joinedload(Order.payments))
        .all()
    )
    cat_of = {pid: cid for pid, cid in db.query(Product.id, Product.category_id).all()}
    cats = {c.id: c for c in db.query(Category).all()}

    daily = {start + timedelta(days=i): {"revenue": 0.0, "profit": 0.0, "orders": 0, "pos": 0.0, "online": 0.0} for i in range(days)}
    hours = [{"hour": h, "revenue": 0.0, "orders": 0} for h in range(24)]
    by_cat: dict = defaultdict(lambda: {"revenue": 0.0, "profit": 0.0, "qty": 0.0})
    by_product: dict = defaultdict(lambda: {"name": "", "qty": 0.0, "revenue": 0.0, "profit": 0.0})
    pay = {"CASH": 0.0, "QR": 0.0, "OTHER": 0.0}
    cur = {"revenue": 0.0, "profit": 0.0, "orders": 0}
    prev = {"revenue": 0.0, "profit": 0.0, "orders": 0}

    for o in orders:
        if not o.completed_at:
            continue
        at = _local(o.completed_at)
        d = at.date()
        total = float(o.total_amount)
        profit = sum(float(i.line_total) - float(i.cost_price) * float(i.quantity) for i in o.items)
        if d < start:
            if d >= prev_start:
                prev["revenue"] += total
                prev["profit"] += profit
                prev["orders"] += 1
            continue
        if d not in daily:
            continue
        row = daily[d]
        row["revenue"] += total
        row["profit"] += profit
        row["orders"] += 1
        row["pos" if o.channel == "POS" else "online"] += total
        cur["revenue"] += total
        cur["profit"] += profit
        cur["orders"] += 1
        hours[at.hour]["revenue"] += total
        hours[at.hour]["orders"] += 1
        for pm in o.payments:
            if pm.status == "SUCCESS":
                key = pm.method if pm.method in ("CASH",) else ("QR" if "QR" in pm.method or pm.method == "BANK" else "OTHER")
                pay[key] += float(pm.amount)
        for i in o.items:
            line_profit = float(i.line_total) - float(i.cost_price) * float(i.quantity)
            c = by_cat[cat_of.get(i.product_id)]
            c["revenue"] += float(i.line_total)
            c["profit"] += line_profit
            c["qty"] += float(i.quantity)
            pr = by_product[i.product_id]
            pr["name"] = i.product_name
            pr["qty"] += float(i.quantity)
            pr["revenue"] += float(i.line_total)
            pr["profit"] += line_profit

    categories = sorted(
        (
            {
                "name": cats[cid].name if cid in cats else "Khác",
                "icon": cats[cid].icon if cid in cats else "📦",
                "revenue": money(v["revenue"]),
                "profit": money(v["profit"]),
                "qty": v["qty"],
            }
            for cid, v in by_cat.items()
        ),
        key=lambda x: x["revenue"],
        reverse=True,
    )
    products = sorted(
        ({"id": pid, **v, "revenue": money(v["revenue"]), "profit": money(v["profit"])} for pid, v in by_product.items()),
        key=lambda x: x["revenue"],
        reverse=True,
    )
    return {
        "days": days,
        "from": start.isoformat(),
        "to": today.isoformat(),
        "totals": {
            **{k: money(v) if k != "orders" else v for k, v in cur.items()},
            "aov": money(cur["revenue"] / cur["orders"]) if cur["orders"] else 0,
            "margin": round(cur["profit"] / cur["revenue"] * 100, 1) if cur["revenue"] else 0,
        },
        "previous": {k: money(v) if k != "orders" else v for k, v in prev.items()},
        "daily": [{"date": k.isoformat(), **{f: money(x) if f != "orders" else x for f, x in v.items()}} for k, v in daily.items()],
        "hours": hours,
        "categories": categories,
        "products": products[:200],
        "payments": {k: money(v) for k, v in pay.items()},
    }


@router.get("/reports/profit")
def profit(db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị xem lợi nhuận")
    from collections import defaultdict

    acc = defaultdict(lambda: {"name": "", "qty": 0, "revenue": 0, "cost": 0, "profit": 0})
    for o in db.query(Order).filter(Order.status == "COMPLETED").options(joinedload(Order.items)).all():
        for i in o.items:
            a = acc[i.product_id]
            a["name"] = i.product_name
            a["qty"] += float(i.quantity)
            a["revenue"] += float(i.line_total)
            a["cost"] += float(i.cost_price) * float(i.quantity)
            a["profit"] = a["revenue"] - a["cost"]
    return sorted(acc.values(), key=lambda x: x["profit"], reverse=True)


@router.get("/reports/inventory-value")
def inventory_value(db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị")
    from app.models import Inventory

    cost = 0.0
    sale = 0.0
    for inv in db.query(Inventory).options(joinedload(Inventory.product)).all():
        if not inv.product or not inv.product.is_active:
            continue
        cost += float(inv.quantity) * float(inv.product.cost_price)
        sale += float(inv.quantity) * float(inv.product.sale_price)
    return {"cost_value": money(cost), "sale_value": money(sale)}


@router.get("/settings")
def get_settings(db: Session = Depends(get_db), _: User = Depends(get_current_staff)):
    return {s.key: s.value for s in db.query(Setting).all()}


class SettingsIn(BaseModel):
    values: dict[str, str]
    # Bắt buộc khi đổi khoá bank.* — tiền khách chuyển khoản đi vào tài khoản này
    pin: str | None = None


class PinIn(BaseModel):
    pin: str


def _check_bank_pin(pin: str | None) -> None:
    import hmac

    if not pin or not hmac.compare_digest(pin, settings.bank_pin):
        raise HTTPException(403, "Sai mã PIN")


@router.post("/settings/unlock")
def unlock_settings(body: PinIn, user: User = Depends(get_current_staff)):
    """Kiểm tra mã PIN trước khi mở phần tài khoản nhận chuyển khoản."""
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị")
    _check_bank_pin(body.pin)
    return {"ok": True}


@router.put("/settings")
def put_settings(body: SettingsIn, db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị")
    if any(k.startswith("bank.") for k in body.values):
        _check_bank_pin(body.pin)
    for k, v in body.values.items():
        row = db.get(Setting, k)
        if row:
            row.value = v
        else:
            db.add(Setting(key=k, value=v))
    return {"ok": True}


@router.get("/promotions")
def promotions(db: Session = Depends(get_db), _: User = Depends(get_current_staff)):
    now = utcnow()
    rows = db.query(Promotion).order_by(Promotion.start_at.desc()).all()
    return [serialize_promo(p, now) for p in rows]


@router.get("/promotions/available")
def available_promotions(db: Session = Depends(get_db), _: User = Depends(get_current_staff)):
    """Mã đang trong hạn cho quầy chọn; mã chưa tới ngày hoặc đã hết hạn không trả về."""
    now = utcnow()
    return [serialize_promo(p, now) for p in live_order_promos(db)]


class PromotionIn(BaseModel):
    code: str
    name: str
    description: str | None = None
    scope: str = "ORDER"
    type: str = "PERCENT"
    value: float
    min_order_amount: float = 0
    max_discount: float | None = None
    usage_limit: int | None = None
    start_date: str
    end_date: str
    is_active: bool = True


def _fill_promo(row: Promotion, body: PromotionIn, db: Session):
    code = body.code.strip().upper().replace(" ", "")
    if not code or len(code) > 30:
        raise HTTPException(400, "Mã phải có 1–30 ký tự")
    dup = db.query(Promotion).filter(Promotion.code == code, Promotion.id != (row.id or 0)).first()
    if dup:
        raise HTTPException(400, f"Mã {code} đã tồn tại")
    if body.scope not in ("ORDER", "NEAR_EXPIRY"):
        raise HTTPException(400, "Phạm vi không hợp lệ")
    if body.type not in ("PERCENT", "AMOUNT"):
        raise HTTPException(400, "Kiểu giảm không hợp lệ")
    if body.scope == "NEAR_EXPIRY" and body.type != "PERCENT":
        raise HTTPException(400, "Giảm cận date phải tính theo %")
    if body.value <= 0 or (body.type == "PERCENT" and body.value > 100):
        raise HTTPException(400, "Mức giảm không hợp lệ")
    start, end = parse_iso_date(body.start_date), parse_iso_date(body.end_date)
    if not start or not end:
        raise HTTPException(400, "Cần ngày bắt đầu và ngày kết thúc")
    if end < start:
        raise HTTPException(400, "Ngày kết thúc phải sau ngày bắt đầu")
    row.code = code
    row.name = body.name.strip() or code
    row.description = (body.description or "").strip() or None
    row.scope = body.scope
    row.type = body.type
    row.value = body.value
    row.min_order_amount = max(0, body.min_order_amount or 0)
    row.max_discount = body.max_discount or None
    row.usage_limit = body.usage_limit or None
    row.start_at = day_start_utc(start)
    row.end_at = day_end_utc(end)
    row.is_active = body.is_active


@router.post("/promotions")
def create_promotion(body: PromotionIn, db: Session = Depends(get_db), _: User = Depends(require("promo.manage"))):
    row = Promotion(used_count=0)
    _fill_promo(row, body, db)
    db.add(row)
    db.flush()
    return serialize_promo(row)


@router.put("/promotions/{promo_id}")
def update_promotion(promo_id: int, body: PromotionIn, db: Session = Depends(get_db), _: User = Depends(require("promo.manage"))):
    row = db.get(Promotion, promo_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy mã")
    _fill_promo(row, body, db)
    db.flush()
    return serialize_promo(row)


@router.delete("/promotions/{promo_id}")
def delete_promotion(promo_id: int, db: Session = Depends(get_db), _: User = Depends(require("promo.manage"))):
    row = db.get(Promotion, promo_id)
    if not row:
        raise HTTPException(404, "Không tìm thấy mã")
    # Mã đã dùng trong hoá đơn thì chỉ tắt, giữ lại để tra cứu
    if row.used_count:
        row.is_active = False
        return {"ok": True, "archived": True}
    db.delete(row)
    return {"ok": True, "archived": False}


@router.get("/audit-logs")
def audit_logs(db: Session = Depends(get_db), user: User = Depends(get_current_staff)):
    if user.role != "ADMIN":
        raise HTTPException(403, "Chỉ quản trị")
    rows = db.query(AuditLog).order_by(AuditLog.id.desc()).limit(100).all()
    return [
        {
            "id": r.id,
            "action": r.action,
            "entity_type": r.entity_type,
            "entity_id": r.entity_id,
            "old_value": r.old_value,
            "new_value": r.new_value,
            "created_at": r.created_at.isoformat() if r.created_at else None,
        }
        for r in rows
    ]
